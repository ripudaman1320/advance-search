#!/usr/bin/env python3
"""
Generate sample DOCX email files for testing the RAG pipeline
Usage: python generate_sample_emails.py
"""

from docx import Document
from docx.shared import Pt, Inches
from datetime import datetime, timedelta
import random
from pathlib import Path

def create_sample_email(subject: str, sender_email: str, sender_name: str, body: str) -> Document:
    """Create a sample email as a DOCX document"""
    doc = Document()
    
    # Add metadata
    doc.core_properties.title = subject
    doc.core_properties.author = sender_name
    
    # Add email header
    doc.add_paragraph(f"From: {sender_name} <{sender_email}>")
    doc.add_paragraph(f"Subject: {subject}")
    doc.add_paragraph(f"Date: {datetime.now().isoformat()}")
    doc.add_paragraph("")
    
    # Add email body
    doc.add_paragraph(body)
    
    # Add signature
    doc.add_paragraph("")
    doc.add_paragraph("--")
    doc.add_paragraph("Best regards,")
    doc.add_paragraph(sender_name)
    
    return doc

def main():
    """Generate sample emails"""
    
    # Create data/raw directory
    output_dir = Path("data/raw")
    output_dir.mkdir(parents=True, exist_ok=True)
    
    # Sample email data
    emails = [
        {
            "subject": "GDG Dev Fest October Volunteering",
            "sender_email": "sawzank@gmail.com",
            "sender_name": "Shasanka Acharya",
            "body": """Hi there!

Thank you for filling out the volunteering form for GDG DevFest.

We're excited to have you on board! 🎉

Please confirm your participation by replying to this email. Once confirmed, 
we'll reach out to you shortly with further details and next steps.

Looking forward to seeing you at the event!

Best regards,
Shasanka Acharya
GDG Sydney"""
        },
        {
            "subject": "Project Implementation Status Update",
            "sender_email": "manager@techcorp.com",
            "sender_name": "Sarah Johnson",
            "body": """Hi team,

I wanted to give you an update on the current project status.

The Q1 implementation is on track. We've completed:
- Backend API development (95% complete)
- Database schema design and setup
- Initial testing framework

Next steps:
- Frontend development (starting next week)
- Integration testing
- Performance optimization

Please make sure to:
1. Complete your assigned tasks by EOW
2. Update the project tracking system daily
3. Attend the standup meetings at 10 AM

We're on schedule to launch by end of Q1. Great work everyone!

Best regards,
Sarah Johnson
Project Manager
TechCorp Inc."""
        },
        {
            "subject": "Re: Machine Learning Models Training",
            "sender_email": "ml-team@research.ai",
            "sender_name": "Dr. Priya Sharma",
            "body": """Hi Ripu,

Thank you for your email regarding the ML model training pipeline.

I've reviewed the implementation and I'm impressed with the architecture. 
Here are my observations:

Strengths:
- Clean separation of concerns between data loading and model training
- Good use of distributed training with PyTorch
- Comprehensive evaluation metrics

Areas for improvement:
- Consider adding cross-validation for better generalization
- Implement early stopping to avoid overfitting
- Add logging for debugging long training runs

I'd like to schedule a meeting to discuss the embedding generation strategy. 
Could you find time next Thursday afternoon?

Please also review the research paper I attached. It covers some recent 
advances in RAG systems that might be relevant to your work.

Best regards,
Dr. Priya Sharma
Lead ML Researcher
Research AI Labs"""
        },
        {
            "subject": "Security Audit Results - Action Required",
            "sender_email": "security@company.com",
            "sender_name": "Alex Chen",
            "body": """Hi Development Team,

We've completed the security audit for the current codebase. 

Here are the critical findings:

CRITICAL (Fix immediately):
1. SQL injection vulnerability in user auth module
   - Severity: High
   - Location: backend/auth/login.py:42
   - Fix: Use parameterized queries

2. Missing CORS validation
   - Severity: High
   - Location: api/middleware.py
   - Fix: Restrict origins to whitelisted domains

HIGH (Fix this week):
3. Weak password hashing
   - Using MD5 instead of bcrypt
   - Update algorithm in user/model.py

4. Exposed API keys in environment
   - Move to Secrets Manager

Please schedule a meeting with your tech lead to discuss remediation. 
All critical issues must be fixed before the next production release.

The full audit report is attached. Let's discuss in our security standup 
on Wednesday at 2 PM.

Best regards,
Alex Chen
Security Lead
Company Inc."""
        }
    ]
    
    # Generate DOCX files
    print("📧 Generating sample email DOCX files...\n")
    
    for i, email_data in enumerate(emails, 1):
        doc = create_sample_email(
            subject=email_data["subject"],
            sender_email=email_data["sender_email"],
            sender_name=email_data["sender_name"],
            body=email_data["body"]
        )
        
        # Create filename from subject
        filename = f"email_{i:02d}_{email_data['subject'].replace(' ', '_')}.docx"
        filepath = output_dir / filename
        
        # Save
        doc.save(str(filepath))
        
        print(f"✓ Created: {filename}")
        print(f"  Subject: {email_data['subject']}")
        print(f"  From: {email_data['sender_email']}")
        print()
    
    print(f"✅ Generated {len(emails)} sample emails in {output_dir}/")
    print("\nThese are ready to ingest into your RAG system!")

if __name__ == "__main__":
    main()
